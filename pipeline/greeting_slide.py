#!/usr/bin/env python3
"""예배 전 화면용 인사말 슬라이드 — 16:9 PPTX, one slide per greeting line (中文 · 拼音 · 한국어)
plus the week's key verse. Same 12192000×6858000 EMU page as the church's own decks, so it can be
dragged straight into the Sunday deck.

Usage: python3 pipeline/greeting_slide.py data/sermons/2026-09-06.json  → data/slides/2026-09-06.greeting.pptx
"""
import json, os, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from enrich import to_pinyin

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NAVY, GOLD, IVORY, MUTED = RGBColor(0x11, 0x18, 0x27), RGBColor(0xE8, 0xB4, 0x4C), RGBColor(0xF7, 0xF3, 0xEA), RGBColor(0xAA, 0xB2, 0xC4)
ZH_FONT, KO_FONT = "Noto Sans CJK SC", "Malgun Gothic"

def box(slide, x, y, w, h, text, size, color, font, bold=False, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return tb

def add_slide(prs, label, zh, py, ko, foot):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = NAVY
    W, H = prs.slide_width, prs.slide_height
    box(s, 0, int(H * 0.10), W, int(H * 0.08), label, 20, GOLD, ZH_FONT, True)
    box(s, int(W * 0.06), int(H * 0.24), int(W * 0.88), int(H * 0.30), zh, 60 if len(zh) <= 14 else 48, IVORY, ZH_FONT, True)
    box(s, int(W * 0.08), int(H * 0.55), int(W * 0.84), int(H * 0.10), py, 26, GOLD, ZH_FONT)
    box(s, int(W * 0.08), int(H * 0.66), int(W * 0.84), int(H * 0.14), ko, 34, MUTED, KO_FONT)
    box(s, int(W * 0.06), int(H * 0.88), int(W * 0.88), int(H * 0.07), foot, 14, MUTED, KO_FONT)

def build(path):
    d = json.load(open(path, encoding="utf-8"))
    prs = Presentation(); prs.slide_width = Emu(12192000); prs.slide_height = Emu(6858000)
    foot = f"{d['church']['zh']} · {d['church']['ko']} · {d['date']} · {d['title']['zh']} / {d['title']['ko']}"
    g = d.get("greeting")
    if g:
        for i, (z, k) in enumerate(zip(g["zh"], g["ko"]), 1):
            add_slide(prs, f"本周问候 · 이번 주 인사 {i}" if len(g["zh"]) > 1 else "本周问候 · 이번 주 인사 — 옆 사람과 서로의 말로", z, to_pinyin(z), k, foot)
    verse = next((b for b in d["video"]["beats"] if b["kind"] == "verse"), None)
    if verse:
        add_slide(prs, f"本周金句 · 이번 주 말씀 — {d['scripture']['ref_zh']} · {d['scripture']['ref_ko']}", verse["zh"], to_pinyin(verse["zh"]), verse["ko"], foot)
    os.makedirs(os.path.join(ROOT, "data", "slides"), exist_ok=True)
    out = os.path.join(ROOT, "data", "slides", f"{d['id']}.greeting.pptx")
    prs.save(out); print("wrote", os.path.relpath(out, ROOT), f"({len(prs.slides)} slides)")

if __name__ == "__main__":
    for p in sys.argv[1:]:
        build(p)
